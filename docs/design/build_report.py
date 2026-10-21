from pathlib import Path
import re
import json
import argparse
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
parser = argparse.ArgumentParser()
parser.add_argument('--source', default='research-content.md')
parser.add_argument('--output', default='炮姐主题模组十种玩法方案调研.docx')
parser.add_argument('--title', default='炮姐主题模组十种玩法方案调研')
parser.add_argument('--header', default='炮姐主题模组  玩法方案调研')
parser.add_argument('--brief', action='store_true')
args = parser.parse_args()
SOURCE = ROOT / args.source
OUTPUT = ROOT / args.output
src = SOURCE.read_text(encoding='utf-8')
pages = src.split('<!-- page -->')
reference_urls = {}
ref_id = None
for line in src.splitlines():
    match = re.match(r'^\[(S\d+)\]', line)
    if match:
        ref_id = match[1]
    elif line.startswith('https://') and ref_id:
        reference_urls.setdefault(ref_id, line.strip())

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)

def font_config(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

normal = doc.styles['Normal']
font_config(normal, 11.5)
normal.paragraph_format.line_spacing = Pt(18.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 25), ('Subtitle', 13), ('Heading 1', 18), ('Heading 2', 13)]:
    sty = doc.styles[name]
    font_config(sty, size, name != 'Subtitle')
    sty.paragraph_format.space_before = Pt(10 if name == 'Heading 2' else 0)
    sty.paragraph_format.space_after = Pt(7 if name == 'Heading 2' else 13)
    sty.paragraph_format.keep_with_next = True
    sty.paragraph_format.line_spacing = 1.15
    ppr = sty.element.get_or_add_pPr()
    for border in list(ppr.findall(qn('w:pBdr'))):
        ppr.remove(border)
font_config(doc.styles['Header'], 9)
font_config(doc.styles['Footer'], 9)

head = sec.header.paragraphs[0]
head.text = args.header
head.paragraph_format.space_after = Pt(0)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('研究稿  2026年10月21日    ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
foot._p.append(field)

def hyperlink(p, label, url, size=None):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    prop = OxmlElement('w:rPr')
    col = OxmlElement('w:color'); col.set(qn('w:val'), '245477'); prop.append(col)
    fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); prop.append(fonts)
    if size:
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); prop.append(sz)
    r.append(prop)
    text = OxmlElement('w:t'); text.text = label; r.append(text)
    link.append(r); p._p.append(link)

def add_text(p, text):
    for part in re.split(r'(\[S\d+\])', text):
        if re.fullmatch(r'\[S\d+\]', part) and part[1:-1] in reference_urls:
            hyperlink(p, part, reference_urls[part[1:-1]])
        else:
            p.add_run(part)

def table(lines, page_num):
    rows = [[s.strip() for s in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r'[:\- ]+', c) for c in row)]
    cols = len(rows[0])
    widths = {
        2: [6.3, 10.7],
        3: [3.1, 5.7, 8.2],
        4: [1.1, 4.3, 6.0, 5.6],
        7: [1.0, 1.5, 1.5, 1.5, 1.5, 1.5, 8.5],
    }[cols]
    if args.brief:
        widths = {2: [4.4, 12.6], 3: [3.4, 6.8, 6.8], 4: [3.1, 4.6, 4.6, 4.7]}[cols]
    elif page_num == 24:
        widths = [3, 2.9, 5.4, 5.7]
    elif page_num == 26:
        widths = [3.0, 5.8, 4.1, 4.1]
    elif page_num == 3:
        widths = [1.0, 1.7, 1.7, 1.7, 1.7, 1.7, 5.5]
    t = doc.add_table(rows=0, cols=cols)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, w in zip(t.columns, widths): c.width = Cm(w)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:' + side)
        for key, val in [('val','single'),('sz','5'),('color','D9D9D9')]:
            el.set(qn('w:' + key), val)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, val in [('top','85'),('bottom','85'),('left','95'),('right','95')]:
        el = OxmlElement('w:' + side); el.set(qn('w:w'), val); el.set(qn('w:type'),'dxa'); margins.append(el)
    t._tbl.tblPr.append(margins)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '30485D' if i == 0 else ('F1F5F8' if i % 2 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = Pt(15.5)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 or (not args.brief and page_num == 3 and j < 6) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            r.font.size = Pt(10)
            r.font.bold = i == 0
            if i == 0: r.font.color.rgb = RGBColor(255,255,255)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.add_run().font.size = Pt(2)

page_stats = []
for idx, page in enumerate(pages, 1):
    if idx > 1:
        doc.add_page_break()
    lines = page.strip().splitlines()
    page_stats.append({'page': idx, 'characters': len(page), 'title': next((x for x in lines if x.startswith('#')), '')})
    k = 0
    while k < len(lines):
        line = lines[k].strip()
        if not line:
            k += 1; continue
        if line.startswith('|'):
            block = []
            while k < len(lines) and lines[k].strip().startswith('|'):
                block.append(lines[k]); k += 1
            table(block, idx)
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            p = doc.add_paragraph(line[3:], 'Heading 1')
        elif line.startswith('### '):
            p = doc.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('https://'):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(9)
            hyperlink(p, '查看来源网页' if 'github.com' not in line else '查看作者源码', line, 10)
        else:
            p = doc.add_paragraph()
            add_text(p, line)
            if idx >= 28:
                p.paragraph_format.line_spacing = Pt(16)
                p.paragraph_format.space_after = Pt(3)
                for r in p.runs: r.font.size = Pt(10)
            elif idx == 1 and not line.startswith('这份') and k < 6:
                p.paragraph_format.space_after = Pt(9)
        k += 1

doc.core_properties.title = args.title
doc.core_properties.subject = '模块职责 独立验收 游戏版本与开发边界' if args.brief else '玩法选择 成长循环 宏观系统 原型验证 版本与联动'
doc.core_properties.author = 'Codex'
doc.core_properties.keywords = 'Minecraft Forge NeoForge 超电磁炮 模组 玩法设计'
settings = doc.settings.element
update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
doc.save(OUTPUT)
plan_name = OUTPUT.stem + '-layout.json' if args.brief else 'layout-plan.json'
(ROOT / plan_name).write_text(json.dumps(page_stats, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'output': str(OUTPUT), 'planned_pages': len(pages), 'characters': len(src), 'sources': len(reference_urls)}, ensure_ascii=False))
